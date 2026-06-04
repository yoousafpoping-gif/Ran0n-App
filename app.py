import streamlit as st
import json
import re
import io
import os
from streamlit_option_menu import option_menu
from gtts import gTTS
from docx import Document
from modules.doc_parser import extract_text
from modules.ai_engine import get_ai_response, get_chat_response
from modules.ui_utils import apply_css

# تغيير اسم البرنامج والأيقونة
st.set_page_config(page_title="Ranon - Your AI Study Partner", page_icon="🧠", layout="wide", initial_sidebar_state="collapsed")
apply_css()

if 'score' not in st.session_state:
    st.session_state['score'] = 0
if 'wrong_answers' not in st.session_state:
    st.session_state['wrong_answers'] = 0
if 'answered_questions' not in st.session_state:
    st.session_state['answered_questions'] = {}
if 'extracted_text' not in st.session_state:
    st.session_state['extracted_text'] = ""

# شريط التنقل بألوان Ranon
mode = option_menu(
    menu_title=None, 
    options=["امتحانات", "ملخصات", "حالات", "فلاش كاردز"],
    icons=["journal-check", "file-text", "clipboard-pulse", "card-text"], 
    default_index=0, 
    orientation="horizontal",
    styles={
        "container": {"padding": "0!important", "background-color": "rgba(42, 59, 111, 0.4)", "border-radius": "15px", "border": "1px solid #6C5CE7", "max-width": "900px", "margin": "0 auto 40px auto"},
        "icon": {"color": "#00D4FF", "font-size": "20px"}, 
        "nav-link": {"font-size": "18px", "text-align": "center", "margin":"0px", "--hover-color": "rgba(108, 92, 231, 0.3)", "color": "#F1F2F6", "font-family": "Cairo"},
        "nav-link-selected": {"background": "linear-gradient(90deg, #6C5CE7, #00D4FF)", "font-weight": "bold", "color": "white"},
    }
)

col_workspace, col_hero = st.columns([1.2, 1], gap="large")

with col_hero: 
    # عرض لوجو Ranon لو الملف موجود
    st.markdown("<div style='text-align: right;'>", unsafe_allow_html=True)
    if os.path.exists("logo.png"):
        st.image("logo.png", width=250)
    else:
        st.markdown("<h1 style='font-size: 4rem; color: #00D4FF; font-family: Poppins, sans-serif; margin-bottom: 0;'>Ranon</h1>", unsafe_allow_html=True)
    
    st.markdown(f"<h3 style='color: #F1F2F6; margin-top: -10px;'>Your AI Study Partner</h3>", unsafe_allow_html=True)
    # استخدام الشعار (Slogan) الخاص بك
    st.markdown("<p style='color: #FFD166; font-size: 1.4rem; font-weight: bold;'>ارفع محاضرتك.. وسيب الباقي علينا.</p></div>", unsafe_allow_html=True)

with col_workspace: 
    st.markdown("<h3 style='text-align: right;'>📂 مساحة العمل</h3>", unsafe_allow_html=True)
    api_key_input = st.text_input("🔑 مفتاح Gemini API:", type="password", placeholder="الصق مفتاحك السري هنا...")
    uploaded_file = st.file_uploader("ارفع المحاضرة (PDF)", type="pdf")
    
    col_diff, col_lang = st.columns(2)
    with col_lang:
        lang = st.radio("🌐 اللغة:", ["العربية", "English"])
    with col_diff:
        difficulty = st.selectbox("🔥 مستوى الصعوبة:", ["أساسيات (مبتدئ)", "متوسط (طالب طب)", "متقدم (زمالة/بورد)"], index=1)
    
    num_q = 10
    if mode in ["امتحانات", "حالات", "فلاش كاردز"]:
        num_q = st.slider("🎯 عدد الأسئلة / الحالات / البطاقات:", min_value=5, max_value=50, value=10, step=5)
        
    col_btn1, col_btn2 = st.columns(2)
    with col_btn1:
        run_button = st.button("🚀 بدء التحليل الذكي")
    with col_btn2:
        if st.button("🔄 إعادة بدء التطبيق"):
            st.session_state.clear()
            st.rerun()

st.markdown("---")

def create_word_doc(quiz_data, title="Ranon Assessment", is_flashcard=False):
    doc = Document()
    doc.add_heading(title, 0)
    if is_flashcard:
        for idx, q in enumerate(quiz_data):
            doc.add_heading(f"Flashcard {idx+1}", level=1)
            doc.add_paragraph(f"Front: {q.get('front')}")
            doc.add_paragraph(f"Back: {q.get('back')}")
            doc.add_paragraph("-" * 20)
    else:
        for idx, q in enumerate(quiz_data):
            doc.add_heading(f"Question {idx+1}", level=1)
            doc.add_paragraph(q.get('question'))
            if q.get('type', 'mcq') == 'mcq':
                for opt in q.get('options', []):
                    doc.add_paragraph(f"- {opt}")
                doc.add_paragraph(f"Correct Answer: {q.get('correct_answer')}")
            else:
                doc.add_paragraph(f"Ideal Answer: {q.get('ideal_answer')}")
            doc.add_paragraph(f"Explanation: {q.get('explanation')}")
            doc.add_paragraph("-" * 20)
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

if run_button:
    if not api_key_input:
        st.error("الرجاء إدخال مفتاح Gemini API.")
    elif not uploaded_file:
        st.error("الرجاء رفع ملف PDF أولاً.")
    else:
        with st.spinner("Ranon يقرأ المحاضرة الآن... ⏳"):
            text = extract_text(uploaded_file)
            if text.startswith("خطأ"):
                st.error(text)
            else:
                st.session_state['extracted_text'] = text
                st.session_state['score'] = 0
                st.session_state['wrong_answers'] = 0
                st.session_state['answered_questions'] = {}
                
                result = get_ai_response(text, mode, api_key_input, num_q, lang, difficulty)
                st.session_state['final_result'] = result
                st.session_state['current_mode'] = mode
                st.session_state['current_lang'] = lang

if 'final_result' in st.session_state:
    result = st.session_state['final_result']
    current_mode = st.session_state['current_mode']
    current_lang = st.session_state.get('current_lang', 'العربية')
    
    st.markdown("## 📊 النتائج المعالجة")
    
    if result.startswith("خطأ"):
        st.error(result)
        
    elif current_mode in ["امتحانات", "حالات"]:
        match = re.search(r'\[.*\]', result, re.DOTALL)
        if match:
            try:
                quiz_data = json.loads(match.group(0))
                total_questions = len(quiz_data)
                
                # لوحة التحكم بألوان Ranon
                st.markdown(f'<div class="card" style="border: 2px solid #00D4FF; background: rgba(42, 59, 111, 0.7);">', unsafe_allow_html=True)
                st.markdown("<h3 style='margin:0; text-align:center; color:#00D4FF;'>📊 لوحة أداء دكتورة رنون</h3>", unsafe_allow_html=True)
                c1, c2, c3 = st.columns(3)
                with c1: st.metric("✅ الإجابات الصحيحة", f"{st.session_state['score']} / {total_questions}")
                with c2: st.metric("❌ الإجابات الخاطئة", f"{st.session_state['wrong_answers']}")
                with c3:
                    percentage = int((st.session_state['score'] / total_questions) * 100) if total_questions > 0 else 0
                    st.metric("🎯 نسبة النجاح الحالية", f"{percentage}%")
                if st.session_state['score'] == total_questions and total_questions > 0:
                    st.balloons()
                    st.success("👑 عبقرية يا رنون! درجة كاملة.. الدكاترة نفسهم ميعرفوش يحلوا كدة!")
                elif percentage >= 80:
                    st.info("✨ الله ينور يا دكتورة، عاش جداً مستواكِ فوق الممتاز!")
                st.markdown('</div>', unsafe_allow_html=True)
                
                word_file = create_word_doc(quiz_data, title=f"Ranon {current_mode}", is_flashcard=False)
                st.download_button(label="📥 تحميل كملف Word", data=word_file, file_name="Ranon_Assessment.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", use_container_width=True)
                
                for idx, q in enumerate(quiz_data):
                    with st.container():
                        st.markdown('<div class="card">', unsafe_allow_html=True)
                        q_type = q.get('type', 'mcq')
                        
                        col_q, col_audio, col_hint = st.columns([4, 0.8, 1])
                        with col_q:
                            st.markdown(f"### ❓ السؤال {idx+1}")
                            st.write(f"**{q.get('question')}**")
                        with col_audio:
                            if st.button("🎧 اسمع", key=f"audio_{idx}_{current_mode}"):
                                tts = gTTS(text=q.get('question'), lang='en' if current_lang == "English" else 'ar', slow=False)
                                audio_fp = io.BytesIO()
                                tts.write_to_fp(audio_fp)
                                audio_fp.seek(0)
                                st.audio(audio_fp, format='audio/mp3')
                        with col_hint:
                            with st.popover("💡 تلميح ذكي"):
                                st.info(q.get('hint', 'فكر في الآلية الفسيولوجية للحالة!'))
                        
                        if q_type == 'mcq':
                            choice = st.radio("اختر الإجابة:", q.get('options', []), key=f"rad_{idx}_{current_mode}", index=None)
                            if st.button("تحقق من إجابتك", key=f"btn_{idx}_{current_mode}"):
                                if choice is None:
                                    st.warning("الرجاء اختيار إجابة أولاً.")
                                else:
                                    is_correct = (choice == q.get('correct_answer'))
                                    prev_status = st.session_state['answered_questions'].get(idx)
                                    if is_correct:
                                        st.success(f"✅ جدعة ي رنون! \n\n **الشرح:** {q.get('explanation')}")
                                        if prev_status != "correct":
                                            if prev_status == "wrong": st.session_state['wrong_answers'] -= 1
                                            st.session_state['score'] += 1
                                            st.session_state['answered_questions'][idx] = "correct"
                                            st.rerun()
                                    else:
                                        st.error(f"❌ حولي تاني ي رنون انتي صح بس الدكاترة مبيحبوش الايجابة دي. \n\n الإجابة الصحيحة: {q.get('correct_answer')}")
                                        if prev_status != "wrong":
                                            if prev_status == "correct": st.session_state['score'] -= 1
                                            st.session_state['wrong_answers'] += 1
                                            st.session_state['answered_questions'][idx] = "wrong"
                                            st.rerun()
                                            
                        elif q_type == 'essay':
                            user_answer = st.text_area("اكتب إجابتك هنا:", key=f"text_{idx}_{current_mode}")
                            if st.button("عرض الإجابة النموذجية", key=f"btn_essay_{idx}_{current_mode}"):
                                st.info(f"💡 **الإجابة النموذجية:** {q.get('ideal_answer')}")
                                st.success(f"**الشرح بالتفصيل:** {q.get('explanation')}")
                        st.markdown('</div>', unsafe_allow_html=True)
            except json.JSONDecodeError:
                st.error("حدث خطأ في قراءة كود الأسئلة.")
                
        else:
            st.error("فشل هيكلة البيانات كـ JSON.")
            
    elif current_mode == "فلاش كاردز":
        match = re.search(r'\[.*\]', result, re.DOTALL)
        if match:
            try:
                flash_data = json.loads(match.group(0))
                word_file = create_word_doc(flash_data, title="Ranon Flashcards", is_flashcard=True)
                st.download_button(label="📥 تحميل الفلاش كاردز كملف Word", data=word_file, file_name="Ranon_Flashcards.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", use_container_width=True)
                
                for idx, card in enumerate(flash_data):
                    with st.container():
                        st.markdown('<div class="card" style="border-right: 5px solid #00D4FF;">', unsafe_allow_html=True)
                        st.markdown(f"### 📇 بطاقة {idx+1}")
                        st.markdown(f"**السؤال:**\n\n {card.get('front')}")
                        
                        show_back = st.checkbox("🔄 اقلب البطاقة", key=f"flash_{idx}")
                        if show_back:
                            st.markdown("<hr style='border-color: #6C5CE7;'>", unsafe_allow_html=True)
                            st.info(f"💡 **الإجابة:**\n\n {card.get('back')}")
                        st.markdown('</div>', unsafe_allow_html=True)
            except json.JSONDecodeError:
                st.error("حدث خطأ في معالجة الفلاش كاردز.")
        else:
            st.error("لم يتم العثور على مصفوفة بطاقات صالحة.")
    else:
        st.markdown(f'<div class="card">{result}</div>', unsafe_allow_html=True)

if st.session_state['extracted_text']:
    st.markdown("---")
    with st.expander("💬 المساعد الطبي لـ Ranon (اسألي في أي نقطة غامضة)"):
        st.markdown("<p style='color: #F1F2F6;'>اكتبي أي جزء غامض أو اطلبي شرح ميكانيزم دوائي معين:</p>", unsafe_allow_html=True)
        user_q = st.text_input("سؤالك الطبي:", key="chat_input_query", placeholder="مثلاً: اشرحيلي ببساطة...")
        if st.button("💬 إرسال"):
            if user_q:
                with st.spinner("Ranon بيجهز الرد... 🩺"):
                    chat_res = get_chat_response(st.session_state['extracted_text'], user_q, api_key_input)
                    st.markdown(f'<div class="card" style="background: rgba(42, 59, 111, 0.8); border-left: 4px solid #00D4FF;">{chat_res}</div>', unsafe_allow_html=True)
            else:
                st.warning("الرجاء كتابة سؤال أولاً.")